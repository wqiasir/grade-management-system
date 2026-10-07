# -*- coding: utf-8 -*-
"""用 PowerPoint 自己生成一页带出场动画的 pptx，导出其 slide1.xml 作为权威参考。

目的：对照我们手写的 <p:timing> 结构，确认 1) 哪些属性不能写；
2) 三层动画树的正确写法；3) TimeLine.MainSequence 的正确读数。
"""

from __future__ import annotations

import os
import zipfile

from pptx import Presentation
from pptx.util import Inches

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "_probe")
os.makedirs(OUT, exist_ok=True)

TEMPLATE = os.path.join(OUT, "pp_template.pptx")
RESULT = os.path.join(OUT, "pp_native_anim.pptx")

#: PowerPoint 动画效果常量：msoAnimEffectFade=10, msoAnimEffectWipe=18(?), 用 Fade 稳妥
MSO_ANIM_EFFECT_FADE = 10
MSO_ANIM_TRIGGER_ON_PAGE_CLICK = 1
MSO_ANIM_TRIGGER_WITH_PREVIOUS = 2
MSO_ANIM_TRIGGER_AFTER_PREVIOUS = 3


def make_template() -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for i, name in enumerate(("形状一", "形状二", "形状三")):
        box = slide.shapes.add_textbox(Inches(1), Inches(1 + i * 1.5), Inches(8), Inches(1))
        box.text_frame.text = name
    prs.save(TEMPLATE)


def animate() -> str:
    import win32com.client as win32

    app = win32.DispatchEx("PowerPoint.Application")
    try:
        pres = app.Presentations.Open(os.path.abspath(TEMPLATE), ReadOnly=False,
                                      Untitled=False, WithWindow=False)
        slide = pres.Slides(1)
        seq = slide.TimeLine.MainSequence
        print("初始 MainSequence:", seq.Count)
        for i in range(1, slide.Shapes.Count + 1):
            eff = seq.AddEffect(
                Shape=slide.Shapes(i),
                EffectId=MSO_ANIM_EFFECT_FADE,
                Trigger=MSO_ANIM_TRIGGER_ON_PAGE_CLICK
                if i == 1 else MSO_ANIM_TRIGGER_AFTER_PREVIOUS,
            )
            print(f"  加入效果 {i}: shape={eff.Shape.Name} effType={eff.EffectType} "
                  f"trigger={eff.Timing.TriggerType}")
        print("加入后 MainSequence:", seq.Count)
        if os.path.exists(RESULT):
            os.remove(RESULT)
        pres.SaveAs(RESULT)
        pres.Close()
    finally:
        try:
            app.Quit()
        except Exception:
            pass
    return RESULT


def dump_timing(path: str) -> None:
    z = zipfile.ZipFile(path)
    names = [n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
    print("\n幻灯片:", names)
    xml = z.read(names[0]).decode("utf-8")
    i = xml.find("<p:timing>")
    print("\n=== PowerPoint 原生 timing（前 4500 字符）===")
    print(xml[i:i + 4500])
    print("\n=== p:seq 元素 ===")
    from lxml import etree
    root = etree.fromstring(z.read(names[0]))
    P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    for s in root.findall(f".//{P}seq"):
        print(etree.tostring(s, pretty_print=True).decode()[:800])


if __name__ == "__main__":
    make_template()
    out = animate()
    dump_timing(out)
